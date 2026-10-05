"""Read issuer-level denial statistics from the CMS Transparency in Coverage PUF (xlsx).

One sheet row is one plan. The issuer-level columns repeat on every plan row of the same
issuer, so the rows are collapsed to one per issuer. A file that fails any quality gate is
rejected outright: nothing is returned.
"""

import logging
from decimal import Decimal
from pathlib import Path
from typing import Annotated, Any, Self

from openpyxl import load_workbook
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    ValidationError,
    field_validator,
    model_validator,
)

from src.db.models import ISSUER_COUNT_COLUMNS, ISSUER_PERCENT_COLUMNS, ExchangeType

logger = logging.getLogger(__name__)

SHEET_NAME = "Transparency 2026 - Ind QHP"
HEADER_ROW = 3
# Legend tokens the source uses in place of a number (not available, suppressed, not required, new).
MISSING_TOKENS = frozenset({"*", "**", "***", "N/A"})
YES_NO = {"Yes": True, "No": False}
MAX_PROBLEMS_IN_MESSAGE = 20

# Row model field -> exact header text in the sheet.
SOURCE_HEADERS = {
    "exchange_type": "Exchange_Type",
    "state": "State",
    "issuer_name": "Issuer_Name",
    "issuer_id": "Issuer_ID",
    "is_new_to_exchange": "Is_Issuer_New_to_Exchange?(Yes_or_No)",
    "claims_received_out_of_network": "Issuer_Claims_Received_Out_of_Network",
    "claims_received_in_network": "Issuer_Claims_Received_In_Network",
    "claims_denied_out_of_network": "Issuer_Claims_Denied_Out_of_Network",
    "claims_denied_in_network": "Issuer_Claims_Denied_In_Network",
    "claims_resubmitted_out_of_network": "Issuer_Claims_Resubmitted_Out_of_Network",
    "claims_resubmitted_in_network": "Issuer_Claims_Resubmitted_In_Network",
    "internal_appeals_filed": "Issuer_Internal_Appeals_Filed",
    "internal_appeals_overturned": "Issuer_Number_Internal_Appeals_Overturned",
    "internal_appeals_overturned_pct": "Issuer_Percent_Internal_Appeals_Overturned",
    "external_appeals_filed": "Issuer_External_Appeals_Filed",
    "external_appeals_overturned": "Issuer_Number_External_Appeals_Overturned",
    "external_appeals_overturned_pct": "Issuer_Percent_External_Appeals_Overturned",
}

Count = Annotated[int | None, Field(ge=0)]
Percent = Annotated[Decimal | None, Field(ge=0, le=100, max_digits=5, decimal_places=2)]


class BatchRejectedError(ValueError):
    """The file failed a quality gate. `problems` lists every reason found."""

    def __init__(self, problems: list[str]) -> None:
        self.problems = problems
        shown = problems[:MAX_PROBLEMS_IN_MESSAGE]
        hidden = len(problems) - len(shown)
        lines = [f"batch rejected, {len(problems)} problem(s):", *shown]
        if hidden:
            lines.append(f"... and {hidden} more")
        super().__init__("\n".join(lines))


class IssuerDenialRow(BaseModel):
    """One validated issuer row, shaped like the `issuer_denial_stats` table."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    plan_year: int = Field(ge=2000, le=2100)
    issuer_id: str = Field(pattern=r"^[0-9]{5}$")
    issuer_name: str = Field(min_length=1, max_length=200)
    state: str = Field(pattern=r"^[A-Z]{2}$")
    exchange_type: ExchangeType
    is_new_to_exchange: bool
    claims_received_out_of_network: Count
    claims_received_in_network: Count
    claims_denied_out_of_network: Count
    claims_denied_in_network: Count
    claims_resubmitted_out_of_network: Count
    claims_resubmitted_in_network: Count
    internal_appeals_filed: Count
    internal_appeals_overturned: Count
    internal_appeals_overturned_pct: Percent
    external_appeals_filed: Count
    external_appeals_overturned: Count
    external_appeals_overturned_pct: Percent

    @field_validator("issuer_id", mode="before")
    @classmethod
    def _issuer_id_as_text(cls, value: Any) -> Any:
        # The sheet stores the id as a number.
        if isinstance(value, int) and not isinstance(value, bool):
            return str(value)
        return value

    @field_validator("is_new_to_exchange", mode="before")
    @classmethod
    def _yes_no_to_bool(cls, value: Any) -> bool:
        if isinstance(value, bool):
            return value
        if value not in YES_NO:
            raise ValueError("must be 'Yes' or 'No'")
        return YES_NO[value]

    @field_validator(*ISSUER_COUNT_COLUMNS, *ISSUER_PERCENT_COLUMNS, mode="before")
    @classmethod
    def _missing_token_to_none(cls, value: Any) -> Any:
        if isinstance(value, str) and value in MISSING_TOKENS:
            return None
        return value

    @field_validator(*ISSUER_PERCENT_COLUMNS, mode="before")
    @classmethod
    def _float_to_exact_decimal(cls, value: Any) -> Any:
        if isinstance(value, float):
            return Decimal(str(value))
        return value

    @model_validator(mode="after")
    def _overturned_not_above_filed(self) -> Self:
        for kind in ("internal", "external"):
            filed = getattr(self, f"{kind}_appeals_filed")
            overturned = getattr(self, f"{kind}_appeals_overturned")
            if filed is not None and overturned is not None and overturned > filed:
                raise ValueError(f"{kind} appeals overturned is greater than {kind} appeals filed")
        return self


def read_issuer_denial_rows(path: Path, plan_year: int) -> list[IssuerDenialRow]:
    """Read the PUF workbook at `path` and return one validated row per issuer.

    The sheet has no year column, so the caller supplies `plan_year`.
    Raises `BatchRejectedError` if any quality gate fails.
    Reading the full file takes a few seconds.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise BatchRejectedError([f"sheet {SHEET_NAME!r} not found"])
        sheet_rows = workbook[SHEET_NAME].iter_rows(min_row=HEADER_ROW, values_only=True)
        header = next(sheet_rows, ())
        positions = {text: index for index, text in enumerate(header)}
        missing = [text for text in SOURCE_HEADERS.values() if text not in positions]
        if missing:
            raise BatchRejectedError(
                [f"row {HEADER_ROW}: missing header {text!r}" for text in missing]
            )

        problems: list[str] = []
        by_issuer: dict[str, IssuerDenialRow] = {}
        for row_number, cells in enumerate(sheet_rows, start=HEADER_ROW + 1):
            if all(cell is None for cell in cells):
                continue  # the sheet ends with a long run of blank rows
            padded = cells + (None,) * (len(header) - len(cells))
            raw = {field: padded[positions[text]] for field, text in SOURCE_HEADERS.items()}
            try:
                row = IssuerDenialRow.model_validate({"plan_year": plan_year, **raw})
            except ValidationError as exc:
                problems.extend(_describe(row_number, exc))
                continue
            earlier = by_issuer.setdefault(row.issuer_id, row)
            if earlier != row:
                problems.append(
                    f"row {row_number}: issuer {row.issuer_id} differs from an earlier row "
                    "of the same issuer"
                )
    finally:
        workbook.close()

    if not problems and not by_issuer:
        problems.append("no data rows")
    if problems:
        raise BatchRejectedError(problems)

    rows = list(by_issuer.values())
    _warn_denied_above_received(rows)
    logger.info("read %d issuer rows for plan year %d", len(rows), plan_year)
    return rows


def _describe(row_number: int, exc: ValidationError) -> list[str]:
    """One line per failed field. Cell values are left out on purpose."""
    return [
        f"row {row_number}: {'.'.join(str(part) for part in error['loc']) or 'row'}: {error['msg']}"
        for error in exc.errors()
    ]


def _warn_denied_above_received(rows: list[IssuerDenialRow]) -> None:
    # Seen in the real file (possibly different counting dates), so this warns instead of rejecting.
    odd = [
        row.issuer_id
        for row in rows
        if row.claims_denied_out_of_network is not None
        and row.claims_received_out_of_network is not None
        and row.claims_denied_out_of_network > row.claims_received_out_of_network
    ]
    if odd:
        logger.warning(
            "%d issuer(s) report more out-of-network claims denied than received: %s",
            len(odd),
            ", ".join(odd),
        )
