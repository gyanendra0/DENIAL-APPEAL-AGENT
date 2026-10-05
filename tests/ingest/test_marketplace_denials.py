"""Loader tests. The fixture workbook is tiny and made up: 3 issuers on 6 plan rows (4 to 9)."""

import logging
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from src.db.models import (
    ISSUER_COUNT_COLUMNS,
    ISSUER_PERCENT_COLUMNS,
    ExchangeType,
    IssuerDenialStats,
)
from src.ingest.marketplace_denials import (
    HEADER_ROW,
    MAX_PROBLEMS_IN_MESSAGE,
    SHEET_NAME,
    BatchRejectedError,
    IssuerDenialRow,
    read_issuer_denial_rows,
    upsert_issuer_denial_rows,
)

CellValue = int | float | str
FIXTURE = Path(__file__).parent / "fixtures" / "tc_puf_sample.xlsx"
PLAN_YEAR = 2026


def _edited_copy(tmp_path: Path, edits: dict[str, CellValue]) -> Path:
    """Copy the fixture with some cells replaced, e.g. {"E4": 123}."""
    workbook = load_workbook(FIXTURE)
    sheet = workbook[SHEET_NAME]
    for cell, value in edits.items():
        sheet[cell] = value
    path = tmp_path / "edited.xlsx"
    workbook.save(path)
    return path


def test_returns_one_row_per_issuer() -> None:
    rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)

    assert [row.issuer_id for row in rows] == ["11111", "22222", "33333"]


def test_reads_every_field_of_an_issuer() -> None:
    first = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)[0]

    assert first == IssuerDenialRow(
        plan_year=PLAN_YEAR,
        issuer_id="11111",
        issuer_name="Sample Health Plan A",
        state="TX",
        exchange_type=ExchangeType.FFE,
        is_new_to_exchange=False,
        claims_received_out_of_network=1000,
        claims_received_in_network=50000,
        claims_denied_out_of_network=300,
        claims_denied_in_network=9000,
        claims_resubmitted_out_of_network=40,
        claims_resubmitted_in_network=500,
        internal_appeals_filed=200,
        internal_appeals_overturned=100,
        internal_appeals_overturned_pct=Decimal("50"),
        external_appeals_filed=None,
        external_appeals_overturned=None,
        external_appeals_overturned_pct=None,
    )


def test_numbers_stored_as_text_become_int() -> None:
    # Cells N4 and Q4 hold the strings "50000" and "40".
    first = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)[0]

    assert first.claims_received_in_network == 50000
    assert first.claims_resubmitted_out_of_network == 40


def test_percent_is_read_as_exact_decimal() -> None:
    second = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)[1]

    assert second.internal_appeals_overturned_pct == Decimal("35.29")


def test_legend_tokens_become_none() -> None:
    new_issuer = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)[2]

    assert new_issuer.is_new_to_exchange is True
    for column in (*ISSUER_COUNT_COLUMNS, *ISSUER_PERCENT_COLUMNS):
        assert getattr(new_issuer, column) is None


def test_skips_blank_trailing_rows() -> None:
    sheet = load_workbook(FIXTURE, read_only=True)[SHEET_NAME]
    row_count = sum(1 for _ in sheet.iter_rows(min_row=HEADER_ROW + 1))
    assert row_count > 6  # the fixture really has blank rows after the 6 plan rows

    assert len(read_issuer_denial_rows(FIXTURE, PLAN_YEAR)) == 3


def test_plan_year_comes_from_the_caller() -> None:
    rows = read_issuer_denial_rows(FIXTURE, 2025)

    assert {row.plan_year for row in rows} == {2025}


@pytest.mark.parametrize(
    ("edits", "expected"),
    [
        ({"E4": 123}, "row 4: issuer_id"),
        ({"C4": "Texas"}, "row 4: state"),
        ({"D4": ""}, "row 4: issuer_name"),
        ({"B4": "XYZ"}, "row 4: exchange_type"),
        ({"F4": "Maybe"}, "row 4: is_new_to_exchange"),
        ({"M4": -1}, "row 4: claims_received_out_of_network"),
        ({"M4": "unknown"}, "row 4: claims_received_out_of_network"),
        ({"M4": 12.5}, "row 4: claims_received_out_of_network"),
        ({"U4": 101}, "row 4: internal_appeals_overturned_pct"),
        ({"U4": 50.123}, "row 4: internal_appeals_overturned_pct"),
        ({"S9": 5, "T9": 6}, "row 9: row: Value error, internal appeals overturned"),
        ({"V9": 5, "W9": 6}, "row 9: row: Value error, external appeals overturned"),
        ({"M5": 999}, "row 5: issuer 11111 differs from an earlier row"),
        ({"E3": "IssuerID"}, "row 3: missing header 'Issuer_ID'"),
    ],
)
def test_rejects_a_bad_file(tmp_path: Path, edits: dict[str, CellValue], expected: str) -> None:
    path = _edited_copy(tmp_path, edits)

    with pytest.raises(BatchRejectedError, match=expected):
        read_issuer_denial_rows(path, PLAN_YEAR)


def test_rejection_lists_every_problem(tmp_path: Path) -> None:
    path = _edited_copy(tmp_path, {"E4": 123, "U7": 101})

    with pytest.raises(BatchRejectedError) as excinfo:
        read_issuer_denial_rows(path, PLAN_YEAR)

    assert len(excinfo.value.problems) == 2


def test_rejection_message_leaves_out_cell_values(tmp_path: Path) -> None:
    path = _edited_copy(tmp_path, {"D4": "S" * 201})

    with pytest.raises(BatchRejectedError) as excinfo:
        read_issuer_denial_rows(path, PLAN_YEAR)

    assert "SSSS" not in str(excinfo.value)


def test_rejects_a_file_without_the_sheet(tmp_path: Path) -> None:
    workbook = load_workbook(FIXTURE)
    workbook[SHEET_NAME].title = "Something else"
    path = tmp_path / "renamed.xlsx"
    workbook.save(path)

    with pytest.raises(BatchRejectedError, match="sheet .* not found"):
        read_issuer_denial_rows(path, PLAN_YEAR)


def test_rejects_a_file_with_no_data_rows(tmp_path: Path) -> None:
    workbook = load_workbook(FIXTURE)
    workbook[SHEET_NAME].delete_rows(HEADER_ROW + 1, 20)
    path = tmp_path / "empty.xlsx"
    workbook.save(path)

    with pytest.raises(BatchRejectedError, match="no data rows"):
        read_issuer_denial_rows(path, PLAN_YEAR)


def test_rejects_an_impossible_plan_year() -> None:
    with pytest.raises(BatchRejectedError, match="plan_year"):
        read_issuer_denial_rows(FIXTURE, 26)


def test_warns_but_loads_when_denied_exceeds_received(
    caplog: pytest.LogCaptureFixture,
) -> None:
    # Issuer 22222 has 150 out-of-network claims denied against 100 received.
    with caplog.at_level(logging.WARNING):
        rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)

    assert len(rows) == 3
    assert "1 issuer(s) report more out-of-network claims denied than received: 22222" in (
        caplog.text
    )


def test_long_rejection_message_is_shortened() -> None:
    problems = [f"row {number}: bad" for number in range(MAX_PROBLEMS_IN_MESSAGE + 5)]

    error = BatchRejectedError(problems)

    assert "... and 5 more" in str(error)
    assert error.problems == problems


def _stored(session: Session) -> list[IssuerDenialStats]:
    """Every stored row, read fresh from the database."""
    session.expire_all()
    query = select(IssuerDenialStats).order_by(
        IssuerDenialStats.plan_year, IssuerDenialStats.issuer_id
    )
    return list(session.scalars(query))


def test_first_load_inserts_one_row_per_issuer(session: Session) -> None:
    rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)

    written = upsert_issuer_denial_rows(session, rows)

    assert written == 3
    assert [stats.issuer_id for stats in _stored(session)] == ["11111", "22222", "33333"]


def test_stores_the_values_that_were_read(session: Session) -> None:
    rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)

    upsert_issuer_denial_rows(session, rows)

    for row, stats in zip(rows, _stored(session), strict=True):
        for field, value in row.model_dump().items():
            assert getattr(stats, field) == value, field


def test_loading_twice_keeps_the_same_rows(session: Session) -> None:
    rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)
    upsert_issuer_denial_rows(session, rows)
    ids_after_first_load = [stats.id for stats in _stored(session)]

    upsert_issuer_denial_rows(session, rows)

    assert [stats.id for stats in _stored(session)] == ids_after_first_load


def test_reload_updates_a_changed_value(session: Session) -> None:
    rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)
    upsert_issuer_denial_rows(session, rows)
    corrected = rows[0].model_copy(
        update={"claims_denied_in_network": 9001, "issuer_name": "Renamed Health Plan"}
    )

    upsert_issuer_denial_rows(session, [corrected])

    first = _stored(session)[0]
    assert first.claims_denied_in_network == 9001
    assert first.issuer_name == "Renamed Health Plan"


def test_reload_advances_updated_at_and_keeps_created_at(session: Session) -> None:
    # now() is fixed for the whole test transaction, so age the rows first.
    long_ago = datetime(2000, 1, 1, tzinfo=UTC)
    rows = read_issuer_denial_rows(FIXTURE, PLAN_YEAR)
    upsert_issuer_denial_rows(session, rows)
    session.execute(update(IssuerDenialStats).values(created_at=long_ago, updated_at=long_ago))

    upsert_issuer_denial_rows(session, rows)

    for stats in _stored(session):
        assert stats.created_at == long_ago
        assert stats.updated_at > long_ago


def test_another_plan_year_adds_new_rows(session: Session) -> None:
    upsert_issuer_denial_rows(session, read_issuer_denial_rows(FIXTURE, 2025))

    upsert_issuer_denial_rows(session, read_issuer_denial_rows(FIXTURE, 2026))

    assert [stats.plan_year for stats in _stored(session)] == [2025] * 3 + [2026] * 3


def test_empty_batch_writes_nothing(session: Session) -> None:
    assert upsert_issuer_denial_rows(session, []) == 0
    assert _stored(session) == []
